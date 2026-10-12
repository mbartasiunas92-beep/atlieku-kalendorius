from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from io import BytesIO
import hashlib
import re
import unicodedata
from urllib.parse import urljoin
from urllib.request import Request, urlopen

import pdfplumber
import streamlit as st
from openpyxl import load_workbook

SOURCE_PAGE = "https://www.nemenkom.lt/buitiniu-ir-pakuociu-atlieku-surinkimo-grafikas/"
MONTHS = {
    "sausis": 1, "vasaris": 2, "kovas": 3, "balandis": 4,
    "geguze": 5, "birzelis": 6, "liepa": 7, "rugpjutis": 8,
    "rugsejis": 9, "spalis": 10, "lapkritis": 11, "gruodis": 12,
}
MONTH_ALIASES = {
    "sausio": 1, "vasario": 2, "kovo": 3, "balandzio": 4,
    "geguzes": 5, "birzelio": 6, "liepos": 7, "rugpjucio": 8,
    "rugsejo": 9, "spalio": 10, "lapkricio": 11, "gruodzio": 12,
}


@dataclass
class Schedule:
    kind: str
    location_text: str
    source_url: str
    events: list[date]


class FileLinks(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.href: str | None = None
        self.label = ""
        self.items: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.label = ""

    def handle_data(self, data: str) -> None:
        if self.href:
            self.label += data

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.href:
            href = urljoin(SOURCE_PAGE, self.href)
            if href.lower().endswith((".pdf", ".xlsx", ".xls")):
                self.items.append((self.label.strip(), href))
            self.href = None


def fetch(url: str) -> bytes:
    request = Request(url, headers={"User-Agent": "WasteCalendarPrototype/0.1"})
    with urlopen(request, timeout=30) as response:
        return response.read()


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.casefold())
    value = "".join(c for c in value if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def latest_links() -> dict[str, tuple[str, str, int]]:
    parser = FileLinks()
    parser.feed(fetch(SOURCE_PAGE).decode("utf-8", errors="replace"))
    found: dict[str, list[tuple[int, str, str, int]]] = {"Buitinės": [], "Pakuotės": [], "Stiklas": []}
    for label, url in parser.items:
        combined = norm(label + " " + url)
        year_match = re.search(r"20\d{2}", label + " " + url)
        year = int(year_match.group()) if year_match else date.today().year
        if "buitiniu" in combined:
            kind = "Buitinės"
        elif "stiklo" in combined:
            kind = "Stiklas"
        elif "pakuociu" in combined:
            kind = "Pakuotės"
        else:
            continue
        period_months = [month for name, month in MONTH_ALIASES.items() if name in combined]
        period = max(period_months, default=0)
        found[kind].append((year * 100 + period, label, url, year))
    selected = {}
    for kind, entries in found.items():
        if entries:
            _, label, url, year = max(entries, key=lambda item: item[0])
            selected[kind] = (label, url, year)
    missing = set(found) - set(selected)
    if missing:
        raise ValueError("Neradau naujausių grafikų: " + ", ".join(sorted(missing)))
    return selected


def dates_in_cell(text: object, year: int, month: int) -> list[date]:
    if text is None:
        return []
    compact = re.sub(r"\s+", "", str(text).casefold())
    days = {int(n) for n in re.findall(r"(?<!\d)(\d{1,2})d?\.?", compact)}
    result = []
    for day in sorted(days):
        try:
            result.append(date(year, month, day))
        except ValueError:
            continue
    return result


def parse_xlsx(blob: bytes, source_url: str, year: int) -> list[Schedule]:
    workbook = load_workbook(BytesIO(blob), data_only=True, read_only=True)
    sheet = workbook.active
    header_row = None
    headers: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [str(value or "").strip() for value in row]
        if "Seniūnija" in values and "Gatvė" in values:
            header_row, headers = row_number, values
            break
    if header_row is None:
        raise ValueError("Buitinių atliekų XLSX faile neatpažinau stulpelių antraščių.")
    indices = {name: headers.index(name) for name in headers if name}
    results: list[Schedule] = []
    inherited = {name: "" for name in ("Seniūnija", "Kaimai")}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        for name in inherited:
            value = row[indices[name]] if name in indices else None
            if value:
                inherited[name] = str(value).strip()
        street = str(row[indices["Gatvė"]] or "").strip() if "Gatvė" in indices else ""
        if not street:
            continue
        location = " · ".join(part for part in (inherited["Seniūnija"], inherited["Kaimai"], street) if part)
        events: list[date] = []
        for header, index in indices.items():
            month = next((n for name, n in MONTHS.items() if norm(header) == name), None)
            if month:
                events.extend(dates_in_cell(row[index], year, month))
        results.append(Schedule("Buitinės", location, source_url, sorted(events)))
    return results


def parse_pdf(blob: bytes, kind: str, source_url: str, year: int) -> list[Schedule]:
    results: list[Schedule] = []
    # Official PDF tables use genitive month labels: Spalio, Lapkričio, Gruodžio.
    month_headers = {"spalio": 10, "lapkricio": 11, "gruodzio": 12}
    with pdfplumber.open(BytesIO(blob)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                header_index = next((i for i, row in enumerate(table) if row and any(norm(str(c or "")) in month_headers for c in row)), None)
                if header_index is None:
                    continue
                headers = [norm(str(c or "")) for c in table[header_index]]
                month_columns = [(i, month_headers[h]) for i, h in enumerate(headers) if h in month_headers]
                for row in table[header_index + 1:]:
                    if not row or len(row) < 3:
                        continue
                    location = str(row[0] or "").strip()
                    if not location or norm(location).startswith("seniunijos pavadinimas"):
                        continue
                    events: list[date] = []
                    for column, month in month_columns:
                        if column < len(row):
                            events.extend(dates_in_cell(row[column], year, month))
                    if events:
                        results.append(Schedule(kind, location, source_url, sorted(set(events))))
    if not results:
        raise ValueError(f"Nepavyko nuskaityti {kind} PDF lentelės.")
    return results


@st.cache_resource(ttl=1800, show_spinner=False)
def load_schedules() -> tuple[list[Schedule], dict[str, tuple[str, str, int]]]:
    links = latest_links()
    rows: list[Schedule] = []
    for kind, (label, url, year) in links.items():
        blob = fetch(url)
        if url.lower().endswith((".xlsx", ".xls")):
            rows.extend(parse_xlsx(blob, url, year))
        else:
            rows.extend(parse_pdf(blob, kind, url, year))
    return rows, links


def search_rows(rows: list[Schedule], query: str) -> list[Schedule]:
    terms = [part for part in norm(query).split() if len(part) >= 3 and not part.isdigit()]
    if not terms:
        return []

    def term_is_present(term: str, text: str) -> bool:
        words = text.split()
        if any(term in word for word in words):
            return True
        # The same place often appears in different Lithuanian grammatical cases
        # across the XLSX and PDFs (e.g. "Didžioji" / "Didžiosios").
        stem = term[:-2] if len(term) >= 7 else term
        return len(stem) >= 5 and any(word.startswith(stem) for word in words)

    return [
        row for row in rows
        if all(term_is_present(term, norm(row.location_text)) for term in terms)
    ]


def ics_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def make_calendar(events: list[tuple[date, Schedule]]) -> bytes:
    now = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Atlieku kalendorius//LT",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:Atliekų išvežimai",
        "BEGIN:VTIMEZONE",
        "TZID:Europe/Vilnius",
        "X-LIC-LOCATION:Europe/Vilnius",
        "BEGIN:DAYLIGHT",
        "TZOFFSETFROM:+0200",
        "TZOFFSETTO:+0300",
        "TZNAME:EEST",
        "DTSTART:20260329T030000",
        "RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=-1SU",
        "END:DAYLIGHT",
        "BEGIN:STANDARD",
        "TZOFFSETFROM:+0300",
        "TZOFFSETTO:+0200",
        "TZNAME:EET",
        "DTSTART:20261025T040000",
        "RRULE:FREQ=YEARLY;BYMONTH=10;BYDAY=-1SU",
        "END:STANDARD",
        "END:VTIMEZONE",
    ]
    seen: set[str] = set()
    for event_date, row in events:
        uid_source = f"{row.kind}|{event_date.isoformat()}|{row.location_text}"
        uid = hashlib.sha256(uid_source.encode("utf-8")).hexdigest() + "@atlieku-kalendorius"
        if uid in seen:
            continue
        seen.add(uid)
        lines.extend([
            "BEGIN:VEVENT",
            f"UID:{uid}",
            f"DTSTAMP:{now}",
            f"DTSTART;TZID=Europe/Vilnius:{event_date:%Y%m%d}T070000",
            f"DTEND;TZID=Europe/Vilnius:{event_date:%Y%m%d}T071500",
            f"SUMMARY:{ics_escape(row.kind + ' atliekų išvežimas')}",
            f"DESCRIPTION:{ics_escape('Grafiko grupė: ' + row.location_text)}",
            "BEGIN:VALARM",
            "ACTION:DISPLAY",
            f"DESCRIPTION:{ics_escape(row.kind + ' atliekų išvežimas po 2 dienų')}",
            "TRIGGER:-P2D",
            "END:VALARM",
            "BEGIN:VALARM",
            "ACTION:DISPLAY",
            f"DESCRIPTION:{ics_escape(row.kind + ' atliekų išvežimas rytoj')}",
            "TRIGGER:-P1D",
            "END:VALARM",
            "END:VEVENT",
        ])
    return ("\r\n".join(lines + ["END:VCALENDAR", ""])).encode("utf-8")


st.set_page_config(page_title="Atliekų kalendorius", page_icon="♻️", layout="centered")
st.title("♻️ Atliekų kalendorius")
st.caption("Vilniaus rajono „Nemenčinės komunalininko“ grafikai pagal jūsų adresą")

with st.expander("Kaip veikia ši versija"):
    st.write("Grafikų nuorodos aptinkamos oficialiame „Nemenkom“ puslapyje, todėl pasikeitus failams jų nereikia įrašyti iš naujo.")
    st.write("PDF grafikuose vienoje eilutėje gali būti kelios gyvenvietės, gatvės ir išimtys. Paieška pateikia atitinkančius grafiko įrašus, kuriuos verta sutikrinti pagal visą adresą.")

try:
    with st.spinner("Gaunu naujausius grafikus iš nemenkom.lt…"):
        schedules, source_links = load_schedules()
except Exception as exc:
    st.error(f"Nepavyko gauti grafikų: {exc}")
    st.link_button("Atidaryti oficialų grafikų puslapį", SOURCE_PAGE)
    st.stop()

st.success("Grafikai įkelti iš oficialaus puslapio.")
for kind, (label, _, _) in source_links.items():
    st.caption(f"{kind}: {label}")

query = st.text_input("Įveskite gyvenvietę ir gatvę", placeholder="Pvz., Didžioji Riešė, Molėtų g.")
if query:
    matches = search_rows(schedules, query)
    if not matches:
        st.info("Pagal šį tekstą įrašo neradau. Pabandykite įvesti trumpesnį vietovės ar gatvės pavadinimą.")
    else:
        today = date.today()
        upcoming_by_kind = {
            kind: sorted(
                ((event, row) for row in matches if row.kind == kind for event in row.events if event >= today),
                key=lambda item: item[0],
            )
            for kind in ("Buitinės", "Pakuotės", "Stiklas")
        }
        calendar_events = [
            (event, row)
            for events in upcoming_by_kind.values()
            for event, row in events
        ]
        if calendar_events:
            st.download_button(
                "Pridėti išvežimus į iPhone kalendorių (.ics)",
                data=make_calendar(calendar_events),
                file_name="atlieku-kalendorius.ics",
                mime="text/calendar",
                help="Failą atidarykite iPhone telefone ir pridėkite į Kalendorių. Įvykiai truks 7:00–7:15, priminimai bus likus 2 dienoms ir 1 dienai.",
            )
        if any(upcoming_by_kind.values()):
            st.subheader("Artimiausi išvežimai")
            for kind, events in upcoming_by_kind.items():
                if not events:
                    continue
                st.markdown(f"### {kind}")
                # Show each category independently so numerous household rows cannot
                # push packaging and glass dates out of the results.
                seen: set[date] = set()
                shown = 0
                for event, row in events:
                    if event in seen:
                        continue
                    seen.add(event)
                    st.markdown(f"**{event.strftime('%d.%m.%Y')}**")
                    st.caption(row.location_text.replace("\n", " "))
                    shown += 1
                    if shown == 3:
                        break
        else:
            st.info("Šiuose grafikuose artėjančių datų neradau.")
        with st.expander(f"Peržiūrėti rastas grafiko grupes ({len(matches)})"):
            for row in matches:
                st.markdown(f"**{row.kind}** — {row.location_text.replace(chr(10), ' ')}")
                st.caption(", ".join(d.strftime("%d.%m.%Y") for d in row.events))

st.divider()
st.caption("Duomenų šaltinis: UAB „Nemenčinės komunalininkas“. Grafikas gali būti pakeistas vežėjo; prieš išstumdami konteinerį sutikrinkite oficialų pranešimą.")
st.link_button("Oficialus grafikų puslapis", SOURCE_PAGE)
